import os
import hmac
import re
import io
import csv
import sqlite3
import json
import time
from pathlib import Path
from contextlib import contextmanager, closing
from flask import Flask, request, jsonify, send_from_directory
from openpyxl import load_workbook
from management import init_management, register_management, capture_open_attendances
from contact_import import phone_key, parse_csv, parse_vcard, inspect_contacts

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "dados"
DATA.mkdir(exist_ok=True)
DB = DATA / "comunicacao.db"
app = Flask(__name__, static_folder="static", static_url_path="/static")
app.json.ensure_ascii = False

@contextmanager
def connect():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    try:
        with con:
            yield con
    finally:
        con.close()

def init_db():
    # One consistent backup before the first additive management migration.
    if DB.exists():
        with closing(sqlite3.connect(DB)) as source:
            old=source.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='wa_conversas'").fetchone()
            migrated=source.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='fluxo_config'").fetchone()
            if old and not migrated:
                backup_dir=DB.parent / 'backups'
                backup_dir.mkdir(exist_ok=True)
                backup_file=backup_dir / ('antes_gestao_'+str(time.time_ns())+'.db')
                with closing(sqlite3.connect(backup_file)) as destination:
                    source.backup(destination)
    with connect() as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS empresas (
          id INTEGER PRIMARY KEY, codigo TEXT NOT NULL UNIQUE,
          nome TEXT NOT NULL, tipo TEXT NOT NULL,
          documento TEXT NOT NULL,
          UNIQUE(tipo, documento)
        );
        CREATE TABLE IF NOT EXISTS contatos (
          id INTEGER PRIMARY KEY, nome TEXT NOT NULL, telefone TEXT NOT NULL UNIQUE
        );
        CREATE TABLE IF NOT EXISTS vinculos (
          empresa_id INTEGER NOT NULL REFERENCES empresas(id) ON DELETE CASCADE,
          contato_id INTEGER NOT NULL REFERENCES contatos(id) ON DELETE CASCADE,
          funcao TEXT NOT NULL DEFAULT '',
          PRIMARY KEY(empresa_id, contato_id)
        );
        CREATE TABLE IF NOT EXISTS wa_conversas (id TEXT PRIMARY KEY, dados TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS wa_mensagens (id TEXT PRIMARY KEY, conversa_id TEXT NOT NULL, timestamp REAL NOT NULL, dados TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS wa_mensagens_conversa ON wa_mensagens(conversa_id,timestamp);
        CREATE TABLE IF NOT EXISTS contatos_excluidos (telefone TEXT PRIMARY KEY);
        CREATE TABLE IF NOT EXISTS wa_omitidas (mensagem_id TEXT PRIMARY KEY, conversa_id TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS wa_midias (mensagem_id TEXT PRIMARY KEY, dados TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS wa_fixadas (mensagem_id TEXT PRIMARY KEY, conversa_id TEXT NOT NULL, expira REAL NOT NULL, duracao INTEGER NOT NULL DEFAULT 0);
        """)
        if "duracao" not in {row[1] for row in db.execute("PRAGMA table_info(wa_fixadas)")}:
            db.execute("ALTER TABLE wa_fixadas ADD COLUMN duracao INTEGER NOT NULL DEFAULT 0")

        init_management(db)

register_management(app, connect)

def digits(s):
    return re.sub(r"\D", "", str(s or ""))

def normalize_company(data):
    codigo = str(data.get("codigo", "")).strip()
    nome = str(data.get("nome", "")).strip()
    tipo = str(data.get("tipo", "")).strip().upper()
    documento = digits(data.get("documento", ""))
    if not codigo or not nome or not documento or tipo not in ("CNPJ", "CPF", "CAEPF", "CEI"):
        raise ValueError("Preencha código, nome, tipo (CNPJ/CPF/CAEPF/CEI) e documento.")
    lengths = {"CNPJ": 14, "CPF": 11, "CAEPF": 14, "CEI": 12}
    if len(documento) != lengths[tipo]:
        raise ValueError(f"{tipo} deve conter {lengths[tipo]} dígitos.")
    return {"codigo":codigo, "nome":nome, "tipo":tipo, "documento":documento}

def company_status(db, item):
    by_code = db.execute("SELECT * FROM empresas WHERE codigo=?", (item["codigo"],)).fetchone()
    by_doc = db.execute("SELECT * FROM empresas WHERE tipo=? AND documento=?", (item["tipo"],item["documento"])).fetchone()
    if (by_code and by_doc and by_code["id"] != by_doc["id"]) or (by_code and (by_code["tipo"],by_code["documento"]) != (item["tipo"],item["documento"])) or (by_doc and by_doc["codigo"] != item["codigo"]):
        return "conflito", "Código ou documento associado a outra empresa."
    if by_code or by_doc:
        return "atualizar", "Cadastro existente: nome será atualizado."
    return "novo", "Novo cadastro."

@app.get("/")
def home():
    return send_from_directory(ROOT / "static", "index.html")

@app.get("/api/empresas")
def empresas():
    with connect() as db:
        return jsonify([dict(r) for r in db.execute("SELECT * FROM empresas ORDER BY nome")])

@app.post("/api/empresas")
def criar_empresa():
    try:
        item = normalize_company(request.get_json(force=True))
        with connect() as db:
            status, info = company_status(db, item)
            if status == "conflito":
                return jsonify(erro=info), 409
            if status == "atualizar":
                return jsonify(erro="Empresa já cadastrada. Utilize a importação para atualizar seu nome."), 409
            db.execute("INSERT INTO empresas(codigo,nome,tipo,documento) VALUES (:codigo,:nome,:tipo,:documento)", item)
        return jsonify(ok=True), 201
    except ValueError as e:
        return jsonify(erro=str(e)), 400

@app.get("/api/contatos")
def contatos():
    selected = request.args.get('telefones')
    candidates = None
    if selected is not None:
        phones = selected.split(',')
        if not 1 <= len(phones) <= 100:
            return jsonify(erro='Selecione de 1 a 100 telefones.'), 400
        candidates = set()
        try:
            for value in phones:
                phone = phone_key('+' + digits(value))
                candidates.update((phone, '+' + phone))
                if phone.startswith('55') and len(phone) in (12, 13):
                    candidates.add(phone[2:])
        except ValueError as exc:
            return jsonify(erro=str(exc)), 400
    with connect() as db:
        if candidates is None:
            query, args = 'SELECT * FROM contatos ORDER BY nome', ()
        else:
            query = 'SELECT * FROM contatos WHERE telefone IN (' + ','.join('?' for _ in candidates) + ') ORDER BY nome'
            args = tuple(candidates)
        rows = [dict(r) for r in db.execute(query, args)]
        for row in rows:
            row["empresas"] = [dict(r) for r in db.execute("""
                SELECT e.id,e.codigo,e.nome,e.tipo,e.documento,v.funcao
                FROM empresas e JOIN vinculos v ON v.empresa_id=e.id
                WHERE v.contato_id=? ORDER BY e.nome""", (row["id"],))]
        return jsonify(rows)

@app.post("/api/contatos")
def criar_contato():
    item = request.get_json(force=True)
    nome = str(item.get("nome","")).strip()
    try:
        telefone = phone_key(item.get("telefone"))
    except ValueError as exc:
        return jsonify(erro=str(exc)), 400
    if not nome or len(nome) > 200 or len(telefone) < 10 or len(telefone) > 15:
        return jsonify(erro="Informe nome com até 200 caracteres e telefone com 10 a 15 dígitos."), 400
    try:
        with connect() as db:
            db.execute("BEGIN IMMEDIATE")
            if any(r['status'] == 'existente' for r in inspect_contacts(db, [{'nome': nome, 'telefone': '+' + telefone}])):
                return jsonify(erro="Telefone já cadastrado. Vincule a empresa ao contato existente."), 409
            db.execute("INSERT INTO contatos(nome,telefone) VALUES (?,?)", (nome,'+' + telefone))
            db.execute('DELETE FROM contatos_excluidos WHERE telefone=?', (telefone,))
        return jsonify(ok=True), 201
    except sqlite3.IntegrityError:
        return jsonify(erro="Telefone já cadastrado. Vincule a empresa ao contato existente."), 409

@app.post("/api/vinculos")
def vincular():
    body = request.get_json(force=True)
    try:
        with connect() as db:
            db.execute("INSERT INTO vinculos(contato_id,empresa_id,funcao) VALUES (?,?,?) ON CONFLICT(contato_id,empresa_id) DO UPDATE SET funcao=excluded.funcao",
                       (int(body["contato_id"]),int(body["empresa_id"]),str(body.get("funcao","")).strip()))
        return jsonify(ok=True)
    except (KeyError,ValueError,sqlite3.IntegrityError):
        return jsonify(erro="Contato ou empresa inválido."), 400

@app.get("/api/pesquisa")
def pesquisa():
    q = request.args.get("q","").strip()
    terms = f"%{q.lower()}%"
    num = f"%{digits(q)}%" if digits(q) else ""
    with connect() as db:
        rows = db.execute("""
          SELECT DISTINCT c.id,c.nome,c.telefone FROM contatos c
          LEFT JOIN vinculos v ON v.contato_id=c.id
          LEFT JOIN empresas e ON e.id=v.empresa_id
          WHERE lower(c.nome) LIKE ? OR c.telefone LIKE ? OR
                lower(e.nome) LIKE ? OR lower(e.codigo) LIKE ? OR e.documento LIKE ?
          ORDER BY c.nome LIMIT 200
        """, (terms,num,terms,terms,num)).fetchall()
        output = []
        for c in rows:
            item = dict(c)
            item["empresas"] = [dict(e) for e in db.execute("""
                SELECT e.codigo,e.nome,e.tipo,e.documento,v.funcao FROM empresas e
                JOIN vinculos v ON v.empresa_id=e.id WHERE v.contato_id=?
                ORDER BY e.nome""",(c["id"],))]
            output.append(item)
        return jsonify(output)

HEADINGS = {
    "codigo": {"codigo","codigo empresa","codigo da empresa","cod empresa","cod"},
    "nome": {"nome","nome da empresa","razao social","empresa"},
    "tipo": {"tipo","tipo inscricao","tipo de inscricao","inscricao tipo"},
    "documento": {"documento","cnpj","cpf","caepf","cei","cnpj cpf caepf cei","cnpj/cpf/caepf/cei","inscricao"}
}
def clean_header(t):
    import unicodedata
    t = unicodedata.normalize("NFKD",str(t or "").lower())
    return re.sub(r"\s+"," ", "".join(x for x in t if not unicodedata.combining(x)).replace("_"," ")).strip()

def parse_excel(blob):
    wb = load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    sheet = wb.active
    values = sheet.iter_rows(values_only=True)
    header = [clean_header(v) for v in next(values,[])]
    cols = {}
    for field, names in HEADINGS.items():
        index = next((i for i,h in enumerate(header) if h in names),None)
        if index is not None:
            cols[field] = index
    if not {"codigo","nome","documento"} <= cols.keys():
        raise ValueError("Planilha precisa ter colunas Código, Nome da empresa e CNPJ/CPF/CAEPF/CEI.")
    records=[]
    for n,row in enumerate(values,2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        def get(k):
            idx=cols.get(k)
            return str(row[idx]).strip() if idx is not None and idx < len(row) and row[idx] is not None else ""
        tipo = get("tipo").upper()
        if not tipo:
            doc_header = header[cols["documento"]]
            tipo = doc_header.upper() if doc_header in ("cnpj","cpf","caepf","cei") else ""
        codigo = get("codigo")
        if codigo.endswith(".0") and codigo[:-2].isdigit(): codigo=codigo[:-2]
        records.append({"linha":n,"codigo":codigo,"nome":get("nome"),"tipo":tipo,"documento":get("documento")})
    wb.close()
    return records

def inspect(records):
    seen_codes,seen_docs=set(),set()
    result=[]
    with connect() as db:
        for record in records:
            n=record["linha"]
            try:
                item=normalize_company(record)
                status,description=company_status(db,item)
                key=(item["tipo"],item["documento"])
                if item["codigo"] in seen_codes or key in seen_docs:
                    status,description="conflito","Registro repetido dentro da planilha."
                seen_codes.add(item["codigo"])
                seen_docs.add(key)
            except ValueError as exc:
                item={k:record.get(k,"") for k in ("codigo","nome","tipo","documento")}
                status,description="invalido",str(exc)
            result.append({"linha":n,**item,"status":status,"detalhe":description})
    return result

@app.post("/api/importar/previa")
def previa():
    uploaded=request.files.get("arquivo")
    if not uploaded or not uploaded.filename.lower().endswith(".xlsx"):
        return jsonify(erro="Envie um arquivo .xlsx."),400
    data=uploaded.read()
    if len(data)>10*1024*1024:
        return jsonify(erro="Arquivo excede 10 MB."),400
    try:
        records=parse_excel(data)
        return jsonify(linhas=inspect(records))
    except Exception as e:
        return jsonify(erro=f"Não foi possível ler a planilha: {e}"),400

@app.post("/api/importar/confirmar")
def confirmar():
    lines=request.get_json(force=True).get("linhas",[])
    if not isinstance(lines,list) or len(lines)>10000:
        return jsonify(erro="Lista inválida."),400
    rows=inspect(lines)
    if any(row["status"] in ("conflito","invalido") for row in rows):
        return jsonify(erro="Resolva as linhas inválidas ou conflitantes antes de confirmar.",linhas=rows),409
    with connect() as db:
        for r in rows:
            if r["status"]=="novo":
                db.execute("INSERT INTO empresas(codigo,nome,tipo,documento) VALUES (?,?,?,?)",
                           (r["codigo"],r["nome"],r["tipo"],r["documento"]))
            else:
                db.execute("UPDATE empresas SET nome=? WHERE codigo=? AND tipo=? AND documento=?",
                           (r["nome"],r["codigo"],r["tipo"],r["documento"]))
    return jsonify(ok=True,importados=len(rows))


@app.post("/api/whatsapp/sync")
def sync_whatsapp_contacts():
    token = os.environ.get("WA_SYNC_TOKEN", "")
    received = request.headers.get("X-Sync-Token", "")
    if not token or not hmac.compare_digest(token, received):
        return jsonify(erro="Não autorizado"), 403
    data = request.get_json(silent=True) or {}
    items = data.get("contatos", [])
    if not isinstance(items, list) or len(items) > 20000:
        return jsonify(erro="Lista inválida"), 400
    novos, existentes, nomes_atualizados = 0, 0, 0
    with connect() as db:
        db.execute('BEGIN IMMEDIATE')
        known = {}
        for row in db.execute('SELECT id,nome,telefone FROM contatos'):
            try:
                known[phone_key(row['telefone'])] = row
            except ValueError:
                continue
        for item in items:
            if not isinstance(item, dict):
                continue
            try:
                phone = phone_key('+' + digits(item.get("telefone")))
            except ValueError:
                continue
            if db.execute('SELECT telefone FROM contatos_excluidos WHERE telefone=?', (phone,)).fetchone():
                continue
            name = str(item.get("nome", "")).strip()[:200]
            if not 10 <= len(phone) <= 15:
                continue
            if phone in known:
                old = known[phone]
                if name and not re.fullmatch(r'[+\d\s().-]+', name) and re.fullmatch(r'[+\d\s().-]+', old['nome']):
                    db.execute('UPDATE contatos SET nome=? WHERE id=?', (name, old['id']))
                    known[phone] = {'id': old['id'], 'nome': name}
                    nomes_atualizados += 1
                existentes += 1
                continue
            if not name:
                name = phone
            inserted = db.execute("INSERT INTO contatos(nome,telefone) VALUES (?,?)", (name,'+' + phone))
            known[phone] = {'id': inserted.lastrowid, 'nome': name}
            novos += 1
    return jsonify(ok=True,novos=novos,existentes=existentes,nomes_atualizados=nomes_atualizados)


@app.post('/api/contatos/importar/previa')
def contact_preview():
    uploaded = request.files.get('arquivo')
    if not uploaded or not uploaded.filename.lower().endswith(('.csv', '.vcf')):
        return jsonify(erro='Selecione um CSV do Google Contacts ou vCard (.vcf).'), 400
    blob = uploaded.read(10 * 1024 * 1024 + 1)
    if len(blob) > 10 * 1024 * 1024:
        return jsonify(erro='Arquivo excede 10 MB.'), 400
    try:
        records = parse_csv(blob) if uploaded.filename.lower().endswith('.csv') else parse_vcard(blob)
        if not records or len(records) > 20000:
            raise ValueError('Arquivo vazio ou excede 20.000 telefones.')
        with connect() as db:
            return jsonify(linhas=inspect_contacts(db, records))
    except (ValueError, UnicodeError, LookupError, csv.Error) as exc:
        return jsonify(erro=f'Não foi possível ler os contatos: {exc}'), 400


@app.post('/api/contatos/importar/confirmar')
def contact_confirm():
    body = request.get_json(silent=True)
    records = body.get('linhas') if isinstance(body, dict) else None
    if not isinstance(records, list) or not records or len(records) > 20000 or any(not isinstance(r, dict) for r in records):
        return jsonify(erro='Lista de contatos inválida.'), 400
    with connect() as db:
        db.execute('BEGIN IMMEDIATE')
        rows = inspect_contacts(db, records)
        if any(r['status'] == 'invalido' for r in rows):
            return jsonify(erro='Corrija os telefones e nomes inválidos antes de importar.', linhas=rows), 409
        new = 0
        for row in rows:
            if row['status'] == 'novo':
                db.execute('INSERT INTO contatos(nome,telefone) VALUES (?,?)', (row['nome'], row['telefone']))
                new += 1
    return jsonify(ok=True, novos=new, preservados=len(rows)-new)

@app.route('/api/mensagens/fixadas', methods=['GET', 'POST', 'DELETE'])
def mensagens_fixadas():
    with connect() as db:
        if request.method == 'GET':
            return jsonify([dict(row) for row in db.execute('SELECT mensagem_id,conversa_id,expira,duracao FROM wa_fixadas WHERE expira>? ORDER BY expira', (time.time(),))])
        data = request.get_json(silent=True)
        if not isinstance(data, dict) or not isinstance(data.get('mensagem_id'), str) or not isinstance(data.get('conversa_id'), str):
            return jsonify(erro='Mensagem inválida.'), 400
        message_id, chat_id = data['mensagem_id'], data['conversa_id']
        if request.method == 'DELETE':
            db.execute('DELETE FROM wa_fixadas WHERE mensagem_id=? AND conversa_id=?', (message_id, chat_id))
            return jsonify(ok=True)
        duration = data.get('duracao')
        if type(duration) is not int or duration not in (86400, 604800, 2592000):
            return jsonify(erro='Escolha 24 horas, 7 dias ou 30 dias.'), 400
        if not db.execute('SELECT id FROM wa_mensagens WHERE id=? AND conversa_id=?', (message_id, chat_id)).fetchone():
            return jsonify(erro='Aguarde a mensagem ser salva e tente novamente.'), 409
        expires = time.time() + duration
        db.execute('INSERT INTO wa_fixadas(mensagem_id,conversa_id,expira,duracao) VALUES (?,?,?,?) ON CONFLICT(mensagem_id) DO UPDATE SET conversa_id=excluded.conversa_id,expira=excluded.expira,duracao=excluded.duracao', (message_id, chat_id, expires, duration))
        return jsonify(mensagem_id=message_id, conversa_id=chat_id, expira=expires, duracao=duration)

@app.route('/api/whatsapp/omitidas', methods=['GET','POST'])
def whatsapp_omitted():
    if not history_authorized():
        return jsonify(erro='Não autorizado'), 403
    with connect() as db:
        if request.method=='GET':
            return jsonify([row[0] for row in db.execute('SELECT mensagem_id FROM wa_omitidas')])
        data=request.get_json(silent=True)
        if not isinstance(data,dict) or not isinstance(data.get('mensagem_id'),str) or not isinstance(data.get('conversa_id'),str):
            return jsonify(erro='Mensagem inválida'),400
        db.execute('INSERT OR IGNORE INTO wa_omitidas VALUES (?,?)',(data['mensagem_id'],data['conversa_id']))
        db.execute('DELETE FROM wa_fixadas WHERE mensagem_id=? AND conversa_id=?',(data['mensagem_id'],data['conversa_id']))
        return jsonify(ok=True)

def history_authorized():
    token = os.environ.get('WA_SYNC_TOKEN', '')
    return bool(token) and hmac.compare_digest(token, request.headers.get('X-Sync-Token', ''))

@app.route('/api/whatsapp/historico', methods=['GET', 'POST'])
def whatsapp_history():
    if not history_authorized():
        return jsonify(erro='Acesso não autorizado'), 403
    if request.method == 'POST':
        payload = request.get_json(silent=True)
        chats = payload.get('conversations') if isinstance(payload, dict) else None
        if not isinstance(chats, list) or len(chats) > 100:
            return jsonify(erro='Histórico inválido'), 400
        for chat in chats:
            if not isinstance(chat, dict) or not isinstance(chat.get('id'), str) or not re.fullmatch(r'(?:\d+@(?:lid|c\.us)|\d+(?:-\d+)?@g\.us)', chat['id']):
                return jsonify(erro='Conversa inválida'), 400
            messages = chat.get('messages', [])
            if not isinstance(messages, list) or any(not isinstance(m, dict) or not isinstance(m.get('id'), str) or not m['id'] or not isinstance(m.get('timestamp'), (int, float)) for m in messages):
                return jsonify(erro='Mensagens inválidas'), 400
    with connect() as db:
        if request.method == 'GET':
            conversations = []
            for row in db.execute('SELECT id,dados FROM wa_conversas'):
                chat = json.loads(row['dados'])
                chat['messages'] = [json.loads(m['dados']) for m in db.execute('SELECT dados FROM wa_mensagens WHERE conversa_id=? AND id NOT IN (SELECT mensagem_id FROM wa_omitidas) ORDER BY timestamp,id', (row['id'],))]
                conversations.append(chat)
            return jsonify(conversations)
        for chat in chats:
            metadata = {k: v for k, v in chat.items() if k != 'messages'}
            db.execute('INSERT INTO wa_conversas(id,dados) VALUES (?,?) ON CONFLICT(id) DO UPDATE SET dados=excluded.dados', (chat['id'], json.dumps(metadata, ensure_ascii=False)))
            for message in chat.get('messages', []):
                db.execute('INSERT INTO wa_mensagens(id,conversa_id,timestamp,dados) VALUES (?,?,?,?) ON CONFLICT(id) DO UPDATE SET dados=excluded.dados,timestamp=excluded.timestamp', (message['id'],chat['id'],message['timestamp'],json.dumps(message,ensure_ascii=False)))
        capture_open_attendances(db)
    return jsonify(ok=True)

@app.route('/api/whatsapp/midia/<path:message_id>', methods=['GET', 'POST'])
def whatsapp_media_cache(message_id):
    if not history_authorized():
        return jsonify(erro='Acesso não autorizado'), 403
    with connect() as db:
        if request.method == 'GET':
            row = db.execute('SELECT dados FROM wa_midias WHERE mensagem_id=?', (message_id,)).fetchone()
            return jsonify(json.loads(row['dados']) if row else None)
        data = request.get_json(silent=True)
        if not isinstance(data,dict) or len(str(data.get('data',''))) > 24*1024*1024:
            return jsonify(erro='Mídia inválida'), 400
        db.execute('INSERT INTO wa_midias(mensagem_id,dados) VALUES (?,?) ON CONFLICT(mensagem_id) DO UPDATE SET dados=excluded.dados', (message_id,json.dumps(data)))
    return jsonify(ok=True)


@app.put('/api/empresas/<int:company_id>')
def editar_empresa(company_id):
    try:
        item = normalize_company(request.get_json(force=True))
        with connect() as db:
            if not db.execute('SELECT id FROM empresas WHERE id=?', (company_id,)).fetchone():
                return jsonify(erro='Empresa não encontrada.'), 404
            if db.execute('SELECT id FROM empresas WHERE id<>? AND (codigo=? OR (tipo=? AND documento=?))', (company_id,item['codigo'],item['tipo'],item['documento'])).fetchone():
                return jsonify(erro='Código ou documento já utilizado por outra empresa.'), 409
            db.execute('UPDATE empresas SET codigo=:codigo,nome=:nome,tipo=:tipo,documento=:documento WHERE id=:id', {**item,'id':company_id})
        return jsonify(ok=True)
    except (ValueError, sqlite3.IntegrityError) as error:
        return jsonify(erro=str(error)), 400

@app.route('/api/contatos/<int:contact_id>', methods=['PUT','DELETE'])
def editar_contato(contact_id):
    with connect() as db:
        if not db.execute('SELECT id FROM contatos WHERE id=?', (contact_id,)).fetchone():
            return jsonify(erro='Contato não encontrado.'), 404
        if request.method == 'DELETE':
            phone=db.execute('SELECT telefone FROM contatos WHERE id=?', (contact_id,)).fetchone()[0]
            db.execute('INSERT OR IGNORE INTO contatos_excluidos(telefone) VALUES (?)', (phone_key(phone),))
            db.execute('DELETE FROM contatos WHERE id=?', (contact_id,))
            return jsonify(ok=True)
        body=request.get_json(silent=True) or {}
        try:
            name=str(body.get('nome','')).strip()
            phone=phone_key(body.get('telefone'))
            if not name or len(name)>200:
                raise ValueError('Informe um nome com até 200 caracteres.')
            existing=[dict(r) for r in db.execute('SELECT nome,telefone FROM contatos WHERE id<>?', (contact_id,))]
            if any(phone_key(r['telefone'])==phone for r in existing):
                return jsonify(erro='Telefone já cadastrado.'), 409
            db.execute('UPDATE contatos SET nome=?,telefone=? WHERE id=?', (name,'+'+phone,contact_id))
            return jsonify(ok=True)
        except (ValueError,sqlite3.IntegrityError) as error:
            return jsonify(erro=str(error)),400

if __name__=="__main__":
    init_db()
    app.run(host="127.0.0.1",port=8765,debug=False)
